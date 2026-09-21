"""Generic-ish parser for the sprint-backlog-style sheets in the workbook.

Columns are located by header keyword rather than fixed position, since the
workbook contains many sheets (one per sprint) whose column layout drifts
slightly release to release.
"""
import openpyxl
import pandas as pd

BLANK = "(blank)"

# canonical field -> keyword(s) to look for in the header cell (lowercased)
_CANON_KEYWORDS = [
    ("seq", ["#"]),
    ("jira_id", ["jira id"]),
    ("project", ["project"]),
    ("jira_story", ["jira story"]),
    ("description", ["description"]),
    ("component", ["component", "componet"]),
    ("assignee", ["assignee"]),
    ("estimated_hours", ["estimat"]),
    ("capitalization", ["capitaliz"]),
    ("priority", ["priority"]),
    ("status", ["status"]),
    ("comments", ["comment"]),
]

_STRONG_HEADER_HINTS = ("project", "assignee", "description", "component", "componet", "status")


def _match_canon(header_text):
    h = (header_text or "").strip().lower()
    if not h:
        return None
    if h == "#":
        return "seq"
    for canon, keywords in _CANON_KEYWORDS:
        if canon == "seq":
            continue
        for kw in keywords:
            if kw in h:
                return canon
    return None


def get_sheet_names(file_like):
    wb = openpyxl.load_workbook(file_like, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _find_header_row(ws, max_scan=10):
    best_row, best_score = 1, -1
    for r in range(1, min(max_scan, ws.max_row) + 1):
        score = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v and str(v).strip().lower() in _STRONG_HEADER_HINTS:
                score += 1
        if score > best_score:
            best_score, best_row = score, r
    return best_row


def parse_sheet(file_like, sheet_name):
    """Returns (dataframe, col_map, header_row)."""
    wb = openpyxl.load_workbook(file_like, data_only=True)
    ws = wb[sheet_name]
    header_row = _find_header_row(ws)

    col_map = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        canon = _match_canon(str(v) if v is not None else "")
        if canon and canon not in col_map:
            col_map[canon] = c

    if "description" not in col_map:
        wb.close()
        return pd.DataFrame(), col_map, header_row

    has_seq_col = "seq" in col_map
    rows = []
    section = None

    def cell(r, canon):
        c = col_map.get(canon)
        return ws.cell(row=r, column=c).value if c else None

    for r in range(header_row + 1, ws.max_row + 1):
        seq = cell(r, "seq")
        description = cell(r, "description")
        project = cell(r, "project")
        component = cell(r, "component")
        assignee = cell(r, "assignee")

        if all(v is None for v in [seq, description, project, component, assignee]):
            continue

        looks_like_divider = description and not project and not component and not assignee

        if has_seq_col:
            if seq is None:
                if looks_like_divider:
                    section = str(description).strip()
                continue
        else:
            if looks_like_divider and str(description).strip().isupper() and len(str(description).split()) <= 3:
                section = str(description).strip()
                continue

        if not description:
            continue

        def norm(v):
            if v is None or str(v).strip() == "":
                return None
            return str(v).strip()

        def norm_num(v):
            if v is None or str(v).strip() == "":
                return None
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v).strip()

        rows.append(
            {
                "row_number": r,
                "section": section,
                "jira_id": norm(cell(r, "jira_id")),
                "project_raw": norm(project),
                "jira_story": norm(cell(r, "jira_story")),
                "description": str(description).strip(),
                "component_raw": norm(component),
                "assignee_raw": norm(assignee),
                "estimated_hours": cell(r, "estimated_hours"),
                "capitalization_raw": norm_num(cell(r, "capitalization")),
                "priority": norm(cell(r, "priority")),
                "status": norm(cell(r, "status")),
                "comments": norm(cell(r, "comments")),
            }
        )

    wb.close()
    df = pd.DataFrame(rows)
    return df, col_map, header_row


def distinct_values(df, column):
    if df.empty or column not in df.columns:
        return []
    vals = df[column].fillna(BLANK).apply(lambda v: BLANK if v == "" else v)
    return sorted(vals.unique().tolist())
